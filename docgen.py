"""Генерация DOCX-документов (python-docx)."""
from datetime import date, datetime
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from calc import fee_order, interest_395, money


def d2s(d) -> str:
    if isinstance(d, str):
        d = date.fromisoformat(d)
    return d.strftime("%d.%m.%Y")


def _new() -> Document:
    doc = Document()
    for s in doc.sections:
        s.left_margin, s.right_margin = Cm(2.5), Cm(1.5)
        s.top_margin = s.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(12)
    return doc


def _p(doc, text="", bold=False, align=None, indent=True, space=6, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(space)
    return p


def _table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = ""
        t.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    return t


def _bytes(doc) -> bytes:
    b = BytesIO()
    doc.save(b)
    return b.getvalue()


def _interest_table(doc, rows, total, principal):
    _table(doc, ["Период", "Дней", "Ставка", "Формула", "Проценты"],
           [(f"{d2s(f)} – {d2s(t)}", n, f"{r:g}%", f"{principal:,.0f} × {n} × {r:g}% / {y}".replace(",", " "),
             money(s)) for f, t, n, r, y, s in rows] + [("ИТОГО", "", "", "", money(total))])


BASIS = {
    "contract": "письменного договора",
    "chat": "договорённости, достигнутой в переписке в мессенджере/по электронной почте",
    "oral": "устной договорённости сторон",
}


def claim_doc(c: dict) -> tuple[bytes, float, float]:
    """c: debtor_name, debtor_addr, my_name, my_addr, my_bank, basis, service, amount, due(ISO)."""
    amount = float(c["amount"])
    due = date.fromisoformat(c["due"])
    today = date.today()
    rows, total = interest_395(amount, due, today)
    doc = _new()
    _p(doc, f"Кому: {c['debtor_name']}\nАдрес: {c['debtor_addr']}", align="right", indent=False)
    _p(doc, f"От: {c['my_name']}\nАдрес: {c['my_addr']}", align="right", indent=False, space=18)
    _p(doc, "ПРЕТЕНЗИЯ", bold=True, align="center", space=0)
    _p(doc, "о погашении задолженности и уплате процентов", bold=True, align="center", space=14)
    _p(doc, f"Между мной и Вами на основании {BASIS.get(c['basis'], BASIS['chat'])} достигнута договорённость "
            f"об оказании услуг (выполнении работ): {c['service']}. Стоимость составила {money(amount)}, "
            f"срок оплаты — до {d2s(due)}.")
    _p(doc, "Свои обязательства я исполнил(а) надлежащим образом, результат передан Вам и принят без замечаний. "
            "Оплата до настоящего времени не поступила.")
    _p(doc, "В соответствии со ст. 309, 310, 711, 781 Гражданского кодекса РФ обязательства должны исполняться "
            "надлежащим образом, односторонний отказ от их исполнения не допускается, а заказчик обязан оплатить "
            "принятый результат. Согласно ст. 395 ГК РФ за пользование чужими денежными средствами вследствие "
            "просрочки их уплаты подлежат уплате проценты по ключевой ставке Банка России.")
    _p(doc, f"Сумма основного долга: {money(amount)}. Проценты по ст. 395 ГК РФ за период с {d2s(due)} "
            f"по {d2s(today)} составляют {money(total)} (расчёт приведён ниже). Проценты продолжают начисляться "
            f"по день фактической уплаты долга.")
    _interest_table(doc, rows, total, amount)
    doc.add_paragraph()
    _p(doc, f"ТРЕБУЮ в течение 10 (десяти) календарных дней с момента получения настоящей претензии выплатить "
            f"{money(amount + total)} (основной долг и проценты на {d2s(today)}) на реквизиты:")
    _p(doc, c["my_bank"], indent=False, italic=True)
    _p(doc, "В случае неисполнения требования в указанный срок я буду вынужден(а) обратиться в суд с требованием "
            "о взыскании долга, процентов, а также судебных расходов, включая государственную пошлину и "
            "расходы на оформление документов, которые будут возложены на Вас.")
    _p(doc, "Приложения: 1) расчёт процентов (в тексте претензии); 2) копии документов, подтверждающих "
            "договорённость и передачу результата (переписка, акт, чек «Мой налог» — при наличии).")
    _p(doc, f"\nДата: {d2s(today)}\t\t\t\t\tПодпись: ____________ / {c['my_name']}", indent=False)
    return _bytes(doc), total, amount + total


def order_doc(c: dict, court: str, arbitr: bool) -> tuple[bytes, float]:
    """Заявление о выдаче судебного приказа (ст. 124 ГПК РФ / ст. 229.4 АПК РФ)."""
    amount = float(c["amount"])
    due = date.fromisoformat(c["due"])
    today = date.today()
    rows, total = interest_395(amount, due, today)
    fee = fee_order(amount, arbitr)
    law = "ст. 229.2–229.4 АПК РФ" if arbitr else "ст. 121–124 ГПК РФ"
    doc = _new()
    _p(doc, f"В {court}", align="right", indent=False)
    _p(doc, f"Взыскатель: {c['my_name']}\nАдрес: {c['my_addr']}\nРеквизиты: {c['my_bank']}",
       align="right", indent=False)
    _p(doc, f"Должник: {c['debtor_name']}\nАдрес: {c['debtor_addr']}", align="right", indent=False)
    _p(doc, f"Сумма требований: {money(amount + total)}\nГоспошлина: {money(fee)}", align="right",
       indent=False, space=14)
    _p(doc, "ЗАЯВЛЕНИЕ", bold=True, align="center", space=0)
    _p(doc, "о вынесении судебного приказа", bold=True, align="center", space=12)
    _p(doc, f"Между взыскателем и должником на основании {BASIS.get(c['basis'], BASIS['chat'])} возникли "
            f"обязательства по оплате услуг (работ): {c['service']}. Стоимость — {money(amount)}, срок оплаты — "
            f"{d2s(due)}. Взыскатель исполнил свои обязательства, должник оплату не произвёл.")
    _p(doc, f"Взыскатель направил должнику претензию, ответа и оплаты не последовало. Требование основано на "
            f"документах, подтверждающих его бесспорность (переписка, акт, чек и др.), и подпадает под "
            f"приказное производство ({law}).")
    _p(doc, f"Проценты по ст. 395 ГК РФ за период с {d2s(due)} по {d2s(today)} составляют {money(total)}:")
    _interest_table(doc, rows, total, amount)
    doc.add_paragraph()
    _p(doc, "На основании изложенного, руководствуясь " + law + ", ПРОШУ:")
    _p(doc, f"1. Вынести судебный приказ о взыскании с должника в пользу взыскателя основного долга "
            f"{money(amount)}, процентов по ст. 395 ГК РФ {money(total)} (с последующим начислением по день "
            f"оплаты), расходов по уплате госпошлины {money(fee)}.")
    _p(doc, "Приложения: квитанция об уплате госпошлины; копия претензии и документ об отправке (опись, чек, "
            "трек-номер); копии документов, подтверждающих требование (переписка, акт, чек НПД); расчёт процентов; "
            "копия заявления для должника (при необходимости).")
    _p(doc, f"\nДата: {d2s(today)}\t\t\t\t\tПодпись: ____________ / {c['my_name']}", indent=False)
    return _bytes(doc), fee


def lawsuit_doc(c: dict, court: str, arbitr: bool) -> bytes:
    """Запасной иск на случай отмены приказа."""
    amount = float(c["amount"])
    due = date.fromisoformat(c["due"])
    today = date.today()
    rows, total = interest_395(amount, due, today)
    from calc import fee_full
    fee = fee_full(amount + total, arbitr)
    doc = _new()
    _p(doc, f"В {court}", align="right", indent=False)
    _p(doc, f"Истец: {c['my_name']}\nАдрес: {c['my_addr']}\nРеквизиты: {c['my_bank']}", align="right", indent=False)
    _p(doc, f"Ответчик: {c['debtor_name']}\nАдрес: {c['debtor_addr']}", align="right", indent=False)
    _p(doc, f"Цена иска: {money(amount + total)}\nГоспошлина (≈): {money(fee)}", align="right", indent=False, space=14)
    _p(doc, "ИСКОВОЕ ЗАЯВЛЕНИЕ", bold=True, align="center", space=0)
    _p(doc, "о взыскании задолженности и процентов", bold=True, align="center", space=12)
    _p(doc, f"На основании {BASIS.get(c['basis'], BASIS['chat'])} истец оказал ответчику услуги (выполнил работы): "
            f"{c['service']}. Стоимость — {money(amount)}, срок оплаты — {d2s(due)}. Результат передан, оплата "
            f"не произведена. Претензия, направленная ответчику, оставлена без удовлетворения. "
            f"Ранее вынесенный судебный приказ отменён в связи с возражениями ответчика.")
    _p(doc, "В силу ст. 309, 310, 711, 781 ГК РФ ответчик обязан оплатить принятый результат; "
            "за просрочку начисляются проценты по ст. 395 ГК РФ.")
    _interest_table(doc, rows, total, amount)
    doc.add_paragraph()
    _p(doc, f"ПРОШУ взыскать с ответчика: основной долг {money(amount)}; проценты по ст. 395 ГК РФ "
            f"{money(total)} с последующим начислением по день оплаты; судебные расходы, включая госпошлину.")
    _p(doc, "Приложения: квитанция об уплате госпошлины (с зачётом ранее уплаченной); претензия и документ "
            "об отправке; доказательства (переписка, акт, чек НПД); определение об отмене приказа; расчёт; "
            "копия иска для ответчика.")
    _p(doc, f"\nДата: {d2s(today)}\t\t\t\t\tПодпись: ____________ / {c['my_name']}", indent=False)
    return _bytes(doc)


def chronology_doc(title: str, events: list[str], my_name: str) -> bytes:
    doc = _new()
    _p(doc, "АКТ-ХРОНОЛОГИЯ СОБЫТИЙ", bold=True, align="center", space=0)
    _p(doc, title, align="center", space=12, indent=False)
    _table(doc, ["№", "Дата", "Событие", "Подтверждение (скриншот №)"],
           [(i + 1, *(e.split("|", 2) + ["", ""])[:2], "") for i, e in enumerate(events)])
    doc.add_paragraph()
    _p(doc, "Скриншоты переписки прилагаются; указана дата сообщения, номер телефона/аккаунт сторон. "
            "Рекомендуется приложить нотариальный протокол осмотра переписки при сумме спора от ~100 тыс. ₽.",
       indent=False, italic=True)
    _p(doc, f"\nДата: {d2s(date.today())}\t\t\tПодпись: ____________ / {my_name}", indent=False)
    return _bytes(doc)


CONTRACT_TYPES = {
    "services": ("Возмездное оказание услуг (самозанятый)", "услуг", "ст. 779–783 ГК РФ", False),
    "repair": ("Подряд (ремонт/строительство)", "работ", "ст. 702–729 ГК РФ", False),
    "creative": ("Авторский заказ (дизайн/тексты/видео)", "творческих работ", "ст. 1288 ГК РФ", True),
    "dev": ("Разработка ПО/сайта", "работ по разработке", "ст. 702, 1296 ГК РФ", True),
    "rent": ("Аренда оборудования", "услуг по аренде", "ст. 606–625 ГК РФ", False),
    "tutor": ("Репетиторство / обучение", "образовательных услуг", "ст. 779 ГК РФ", False),
    "beauty": ("Бьюти-услуги", "услуг", "ст. 779 ГК РФ", False),
    "smm": ("SMM / маркетинг", "услуг", "ст. 779 ГК РФ", True),
}


def contract_doc(kind: str, c: dict) -> bytes:
    title, noun, law, ip = CONTRACT_TYPES[kind]
    pre = int(c["prepay"])
    doc = _new()
    _p(doc, f"ДОГОВОР № ___ ({title.lower()})", bold=True, align="center", space=0)
    _p(doc, f"Дата: {d2s(date.today())}", align="center", space=12, indent=False)
    _p(doc, f"{c['exec_name']} (далее — «Исполнитель»), и {c['cust_name']} (далее — «Заказчик»), "
            f"заключили настоящий договор ({law}):")
    items = [
        ("1. Предмет", f"Исполнитель обязуется оказать Заказчику ряд {noun}: {c['subject']}. "
                       f"Состав, объём и сроки — в п. 2 и приложениях/ТЗ, согласованных сторонами в письменной форме."),
        ("2. Сроки и приёмка", f"Срок выполнения: {c['term']}. Результат передаётся поэтапно по акту; Заказчик обязан "
                               f"принять результат или направить мотивированные замечания в течение 3 рабочих дней, "
                               f"иначе результат считается принятым."),
        ("3. Цена и оплата", f"Стоимость: {money(float(c['price']))}. Предоплата {pre}% — до начала работ; "
                             f"остаток — в течение 3 банковских дней после подписания акта. Исполнитель — плательщик НПД, "
                             f"чек формируется в «Мой налог»."),
        ("4. Неустойка", "За просрочку оплаты Заказчик уплачивает неустойку 0,1% от суммы долга за каждый день "
                         "просрочки (ст. 330 ГК РФ)."),
        ("5. Удержание результата", "До полной оплаты Исполнитель вправе удерживать результат работ (ст. 712 ГК РФ) "
                                    "и приостановить дальнейшее выполнение."),
    ]
    if ip:
        items.append(("6. Права на результат", "Исключительные права на результат переходят к Заказчику только после "
                                               "полной оплаты. До этого момента использование результата не допускается; "
                                               "неоплаченное использование является нарушением прав Исполнителя."))
    n = len(items) + 1
    items += [
        (f"{n}. Переписка", f"Стороны признают юридическую силу переписки в мессенджерах и по e-mail по реквизитам: "
                            f"Исполнитель — {c['exec_contact']}; Заказчик — {c['cust_contact']}. Сообщения с этих "
                            f"аккаунтов считаются исходящими от стороны."),
        (f"{n + 1}. Споры", "Претензионный срок — 10 календарных дней. Споры рассматриваются по месту жительства/"
                            "нахождения Исполнителя."),
    ]
    for h, t in items:
        _p(doc, h, bold=True, indent=False, space=2)
        _p(doc, t)
    _p(doc, f"\nИсполнитель: {c['exec_name']}\t\t\tЗаказчик: {c['cust_name']}\n"
            f"Подпись: ____________\t\t\t\tПодпись: ____________", indent=False)
    return _bytes(doc)


def act_doc(c: dict) -> bytes:
    doc = _new()
    _p(doc, "АКТ СДАЧИ-ПРИЁМКИ", bold=True, align="center", space=0)
    _p(doc, f"к договору от {c['contract_date']} • {d2s(date.today())}", align="center", space=12, indent=False)
    _p(doc, f"Исполнитель {c['exec_name']} передал, а Заказчик {c['cust_name']} принял результат: {c['subject']}. "
            f"Стоимость — {money(float(c['price']))}. Работы выполнены в полном объёме, в срок; "
            f"претензий по объёму, качеству и срокам Заказчик не имеет.")
    _p(doc, f"\nИсполнитель: ____________ / {c['exec_name']}\nЗаказчик: ____________ / {c['cust_name']}", indent=False)
    return _bytes(doc)
